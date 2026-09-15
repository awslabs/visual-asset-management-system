# Copyright 2026 Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: Apache-2.0

"""Office documents (.docx .pptx .xlsx): text and document properties through python-docx, python-pptx and
openpyxl -> sys_document / sys_data, the text excerpt, and the full text with slide or sheet offsets for content
chunking. Nothing is rasterised: the analysis model works from the text."""

import csv
import datetime
import io
import re
import zipfile
from typing import Dict, List, Optional, Tuple

import docx
import openpyxl
from defusedxml import ElementTree
from defusedxml.ElementTree import ParseError
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError

from contentChunks import CONTENT_TEXT_MAX_CHARS

from .common import (
    CLASS_DATA,
    CLASS_DOCUMENT,
    DATA_CELL_MAX_CHARS,
    DATA_MAX_COLUMNS,
    DATA_SAMPLE_ROWS,
    BranchResult,
    ExtractContext,
    human_count,
    truncate_text,
)

_WORD = re.compile(r"\w+", re.UNICODE)
_PROPERTY_MAX_CHARS = 500
_SHEET_NAMES_MAX = 50
# Rows counted per sheet; the count stops here so a million-row workbook is not read to its end for a number.
XLSX_MAX_ROWS_COUNTED = 200_000
# Slides and sheets of the captured full text are joined with a blank line; page offsets point at each part's
# first character.
_PAGE_SEPARATOR = "\n\n"
# The extended-properties part carries the page count the authoring application recorded when it saved the file.
_APP_PROPERTIES_PART = "docProps/app.xml"
_APP_PROPERTIES_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
# What the three libraries raise for a file that is not a readable package.
_OPEN_ERRORS = (DocxPackageNotFoundError, PptxPackageNotFoundError, InvalidFileException, zipfile.BadZipFile,
                KeyError, ValueError, OSError)


class _TextCollector:
    """Assembles the captured full text part by part under CONTENT_TEXT_MAX_CHARS, recording each part's page
    entry; a part that does not fit is cut and every later part is dropped."""

    def __init__(self):
        self.parts: List[str] = []
        self.offsets: List[dict] = []
        self.length = 0
        self.truncated = False

    def add(self, text: str, page: int, name: Optional[str] = None) -> None:
        if self.truncated:
            return
        start = self.length + (len(_PAGE_SEPARATOR) if self.parts else 0)
        room = CONTENT_TEXT_MAX_CHARS - start
        if room <= 0:
            self.truncated = True
            return
        if len(text) > room:
            text, self.truncated = text[:room], True
        entry = {"page": page, "start": start}
        if name:
            entry["name"] = name
        self.offsets.append(entry)
        self.parts.append(text)
        self.length = start + len(text)

    @property
    def text(self) -> str:
        return _PAGE_SEPARATOR.join(self.parts)


def _iso(value) -> Optional[str]:
    if not isinstance(value, datetime.datetime):
        return None
    if value.tzinfo is not None:
        value = value.astimezone(datetime.timezone.utc).replace(tzinfo=None)
    return value.strftime("%Y-%m-%dT%H:%M:%SZ")


def _core_properties(core) -> Dict[str, object]:
    """title, author, subject, keywords, createdAt and modifiedAt from an OPC core-properties part."""
    properties: Dict[str, object] = {}
    for key in ("title", "author", "subject", "keywords"):
        value = getattr(core, key, None)
        if value:
            properties[key] = str(value)[:_PROPERTY_MAX_CHARS]
    created = _iso(getattr(core, "created", None))
    if created:
        properties["createdAt"] = created
    modified = _iso(getattr(core, "modified", None))
    if modified:
        properties["modifiedAt"] = modified
    return properties


def app_page_count(path: str) -> Optional[int]:
    """The ``Pages`` count of the package's extended properties, or None when the package records none."""
    try:
        with zipfile.ZipFile(path) as package:
            if _APP_PROPERTIES_PART not in package.namelist():
                return None
            root = ElementTree.fromstring(package.read(_APP_PROPERTIES_PART))
    except (zipfile.BadZipFile, ParseError, KeyError, OSError):
        return None
    node = root.find(f"{_APP_PROPERTIES_NS}Pages")
    try:
        return int(node.text) if node is not None and node.text else None
    except ValueError:
        return None


def _document_facts(result: BranchResult, section: Dict[str, object]) -> None:
    for key in ("title", "author", "subject"):
        if section.get(key):
            result.facts[key] = str(section[key])
    if section.get("createdAt"):
        result.facts["created"] = str(section["createdAt"])


def _unreadable(result: BranchResult, section_key: str, fmt: str, exc: Exception) -> BranchResult:
    result.attributes[section_key] = {"format": fmt, "decodable": False}
    result.warnings.append(f"{fmt} could not be opened: {exc}")
    return result


def _docx_text(document) -> Tuple[str, bool]:
    """Paragraph texts and table rows (cells tab-separated) in document order, newline-joined and collected
    only up to CONTENT_TEXT_MAX_CHARS: the line that crosses the cap is cut there, nothing after it is read,
    and the second value records that cut."""
    lines: List[str] = []
    length = 0
    truncated = False
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            block = [Paragraph(child, document).text]
        elif tag == "tbl":
            block = ["\t".join(cell.text.strip() for cell in row.cells) for row in Table(child, document).rows]
        else:
            continue
        for line in block:
            if not lines and not line:
                continue
            room = CONTENT_TEXT_MAX_CHARS - length - (1 if lines else 0)
            if room <= 0:
                truncated = True
                break
            if len(line) > room:
                line, truncated = line[:room], True
            length += len(line) + (1 if lines else 0)
            lines.append(line)
        if truncated:
            break
    return "\n".join(lines).strip("\n"), truncated


def extract_docx(path: str, ctx: ExtractContext) -> BranchResult:
    result = BranchResult(file_class=CLASS_DOCUMENT)
    try:
        document = docx.Document(path)
    except _OPEN_ERRORS as exc:
        return _unreadable(result, "sys_document", "DOCX", exc)
    text, truncated = _docx_text(document)
    words = len(_WORD.findall(text))
    sys_document: Dict[str, object] = {"format": "DOCX", "decodable": True, "wordCount": words, "hasText": words > 0}
    pages = app_page_count(path)
    if pages:
        sys_document["pageCount"] = pages
    sys_document.update(_core_properties(document.core_properties))
    result.attributes["sys_document"] = sys_document
    if pages:
        result.facts["pages"] = human_count(pages, "page")
    result.facts["words"] = human_count(words, "word")
    _document_facts(result, sys_document)
    result.text_excerpt = truncate_text(text, ctx.max_text_chars)
    if ctx.capture_full_text:
        # Word Open XML records page breaks only as rendering hints, so the full text carries no page entries.
        result.full_text = text
        result.full_text_truncated = truncated
    return result


def _slide_text(slide) -> str:
    lines: List[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            frame = shape.text_frame.text
            if frame.strip():
                lines.append(frame)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                lines.append("\t".join(cell.text.strip() for cell in row.cells))
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame
        if notes is not None and notes.text.strip():
            lines.append("Notes: " + notes.text)
    return "\n".join(lines)


def extract_pptx(path: str, ctx: ExtractContext) -> BranchResult:
    result = BranchResult(file_class=CLASS_DOCUMENT)
    try:
        presentation = Presentation(path)
    except _OPEN_ERRORS as exc:
        return _unreadable(result, "sys_document", "PPTX", exc)
    slide_count = len(presentation.slides)
    # Slides are collected in order until the cap; a slide beyond it is never read.
    collector = _TextCollector()
    for number, slide in enumerate(presentation.slides, start=1):
        collector.add(_slide_text(slide), number)
        if collector.truncated:
            break
    text = collector.text
    words = len(_WORD.findall(text))
    sys_document: Dict[str, object] = {
        "format": "PPTX", "decodable": True, "slideCount": slide_count, "pageCount": slide_count,
        "wordCount": words, "hasText": words > 0}
    sys_document.update(_core_properties(presentation.core_properties))
    result.attributes["sys_document"] = sys_document
    result.facts["slides"] = human_count(slide_count, "slide")
    result.facts["words"] = human_count(words, "word")
    _document_facts(result, sys_document)
    result.text_excerpt = truncate_text(text, ctx.max_text_chars)
    if ctx.capture_full_text:
        result.full_text = text
        result.page_offsets = collector.offsets
        result.full_text_truncated = collector.truncated
    return result


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.isoformat()
    return str(value)[:DATA_CELL_MAX_CHARS]


def extract_xlsx(path: str, ctx: ExtractContext) -> BranchResult:
    result = BranchResult(file_class=CLASS_DATA)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except _OPEN_ERRORS as exc:
        return _unreadable(result, "sys_data", "XLSX", exc)
    collector = _TextCollector()
    first_columns: List[str] = []
    first_sample: List[List[str]] = []
    first_rows = 0
    first_capped = False
    try:
        sheet_names = list(workbook.sheetnames)
        for number, name in enumerate(sheet_names, start=1):
            header: Optional[List[str]] = None
            sample: List[List[str]] = []
            rows = 0
            capped = False
            lines = [f"Sheet: {name}"]
            line_chars = len(lines[0])
            # The text this sheet may still add: the cap less what the earlier sheets used.
            want_text = ctx.capture_full_text and not collector.truncated
            text_budget = CONTENT_TEXT_MAX_CHARS - collector.length
            for row in workbook[name].iter_rows(values_only=True):
                cells = [_cell(value) for value in row[:DATA_MAX_COLUMNS]]
                line = "\t".join(cells)
                if header is None:
                    header = cells
                    lines.append(line)
                    line_chars += len(line) + 1
                    continue
                if rows < XLSX_MAX_ROWS_COUNTED:
                    rows += 1
                else:
                    capped = True
                if len(sample) < DATA_SAMPLE_ROWS:
                    sample.append(cells)
                if want_text and line_chars < text_budget:
                    lines.append(line)
                    line_chars += len(line) + 1
                # One pass feeds the sample, the count and the text; it ends once none of them needs another row.
                if capped and len(sample) >= DATA_SAMPLE_ROWS and (not want_text or line_chars >= text_budget):
                    break
            if number == 1:
                first_columns, first_sample, first_rows, first_capped = header or [], sample, rows, capped
            if want_text:
                collector.add("\n".join(lines), number, name)
    finally:
        workbook.close()
    sys_data: Dict[str, object] = {
        "format": "XLSX",
        "decodable": True,
        "sheetCount": len(sheet_names),
        "sheetNames": sheet_names[:_SHEET_NAMES_MAX],
        "columnCount": len(first_columns),
        "columns": first_columns,
        "rowCount": first_rows,
        "rowCountExact": not first_capped,
        "sampleRows": first_sample,
    }
    result.attributes["sys_data"] = sys_data
    result.facts["sheets"] = human_count(len(sheet_names), "sheet")
    result.facts["columns"] = human_count(len(first_columns), "column")
    result.facts["rows"] = human_count(first_rows, "row") + (" (counting stopped)" if first_capped else "")
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(first_columns)
    writer.writerows(first_sample)
    result.text_excerpt = truncate_text(buffer.getvalue(), ctx.max_text_chars)
    if ctx.capture_full_text:
        result.full_text = collector.text
        result.page_offsets = collector.offsets
        result.full_text_truncated = collector.truncated
    return result
